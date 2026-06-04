import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
import io
from config import config


def export_to_excel_combined(entries_data, filename="cbt_diary.xlsx"):
    """
    Экспорт в одну таблицу: все данные в одном листе с эмоциями в строке
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "КПТ Дневник"
    
    # ШАПКА таблицы
    headers = [
        "Дата",
        "Ситуация (кратко)",
        "Эмоции (по шкале 0-100)",
        "Автоматические мысли",
        "Разумная реакция",
        "Результат"
    ]
    
    # Записываем заголовки со стилями
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Записываем данные
    if entries_data:
        for row_idx, entry in enumerate(entries_data, start=2):
            # Преобразуем entry в словарь
            if hasattr(entry, '__dict__'):
                entry_dict = entry.__dict__
                # Убираем служебные поля SQLAlchemy
                entry_dict = {k: v for k, v in entry_dict.items() 
                            if not k.startswith('_')}
            elif isinstance(entry, dict):
                entry_dict = entry
            else:
                print(f"⚠️ Неизвестный формат записи: {type(entry)}")
                continue
            
            # 1. ДАТА
            created_at = entry_dict.get('created_at')
            if isinstance(created_at, datetime):
                date_str = created_at.strftime("%d.%m.%Y %H:%M")
            elif created_at:
                date_str = str(created_at)[:16]
            else:
                date_str = ""
            ws.cell(row=row_idx, column=1, value=date_str)
            
            # 2. СИТУАЦИЯ (РАСШИФРОВЫВАЕМ)
            situation = entry_dict.get('situation', '')
            if situation:
                situation = config.decrypt_text(str(situation))
                ws.cell(row=row_idx, column=2, value=situation)
            else:
                ws.cell(row=row_idx, column=2, value="")
            
            # 3. ЭМОЦИИ
            emotions_text = ""
            emotions = entry_dict.get('emotions', {})
            
            if isinstance(emotions, dict) and emotions:
                emotions_list = []
                for emotion_name, emotion_data in emotions.items():
                    if isinstance(emotion_data, dict):
                        intensity = emotion_data.get('intensity', '')
                        reassessment = emotion_data.get('reassessment', '')
                        
                        if intensity is not None and reassessment is not None:
                            emotions_list.append(f"{emotion_name}: {intensity}%→{reassessment}%")
                        elif intensity is not None:
                            emotions_list.append(f"{emotion_name}: {intensity}%")
                        else:
                            emotions_list.append(emotion_name)
                    elif isinstance(emotion_data, (int, float)):
                        emotions_list.append(f"{emotion_name}: {emotion_data}%")
                    else:
                        emotions_list.append(emotion_name)
                
                emotions_text = ", ".join(emotions_list)
            elif isinstance(emotions, list):
                emotions_text = ", ".join(str(e) for e in emotions)
            
            ws.cell(row=row_idx, column=3, value=emotions_text)
            
            # 4. АВТОМАТИЧЕСКИЕ МЫСЛИ (РАСШИФРОВЫВАЕМ)
            thought = entry_dict.get('automatic_thought', '')
            if thought:
                thought = config.decrypt_text(str(thought))
                ws.cell(row=row_idx, column=4, value=thought)  # ПОЛНЫЙ ТЕКСТ
            else:
                ws.cell(row=row_idx, column=4, value="")

            # 5. РАЗУМНАЯ РЕАКЦИЯ (РАСШИФРОВЫВАЕМ)
            response = entry_dict.get('rational_response', '')
            if response:
                response = config.decrypt_text(str(response))
                ws.cell(row=row_idx, column=5, value=response)  # ПОЛНЫЙ ТЕКСТ
            else:
                ws.cell(row=row_idx, column=5, value="")
            
            # 6. РЕЗУЛЬТАТ
            result = entry_dict.get('result', '')
            if result:
                result = config.decrypt_text(str(result))
                ws.cell(row=row_idx, column=6, value=result)
            else:
                # Формируем результат из эмоций если нет прямого результата
                if isinstance(emotions, dict):
                    changes = []
                    for emotion_name, emotion_data in emotions.items():
                        if isinstance(emotion_data, dict):
                            intensity = emotion_data.get('intensity')
                            reassessment = emotion_data.get('reassessment')
                            
                            if intensity is not None and reassessment is not None:
                                change = reassessment - intensity
                                if change > 0:
                                    changes.append(f"{emotion_name} ↑{change}%")
                                elif change < 0:
                                    changes.append(f"{emotion_name} ↓{abs(change)}%")
                                else:
                                    changes.append(f"{emotion_name} →")
                    
                    if changes:
                        result = ", ".join(changes)
                    else:
                        result = "Без изменений"
                
                ws.cell(row=row_idx, column=6, value=result)
    
    # Настраиваем ширину колонок
    column_widths = [15, 30, 40, 40, 40, 30]
    
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    # Включаем перенос текста для всех ячеек
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    # Сохраняем в байтовый поток
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    return excel_buffer